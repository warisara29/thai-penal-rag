"""Judge validation — human vs LLM-judge agreement (design §5/§10).

Two modes:
  --sample : draw a stratified, BLIND audit sheet for a human to label.
             Writes runs/<run>/judge_audit/audit_sheet_blind.csv  (no judge columns —
             so the annotator is not anchored) and audit_key.csv (hidden judge verdicts).
  --score  : after the human fills HUMAN_correct (1/0) and HUMAN_score (1-5) in the
             blind sheet, join with the key and report Cohen's kappa (binary correct),
             raw agreement, Pearson + Spearman on the 1-5 score, and the confusion matrix.

Stratified across arms of differing quality (spanning the answer-quality range) and
balanced across the judge's correct/incorrect classes so kappa is well-conditioned.
Deterministic (fixed seed). Pure Python.

Usage:
  .venv/bin/python -m analysis.judge_audit --run runs/2026-08-07 --sample --n 80
  # ... human fills HUMAN_correct / HUMAN_score in audit_sheet_blind.csv ...
  .venv/bin/python -m analysis.judge_audit --run runs/2026-08-07 --score
"""

from __future__ import annotations

import argparse
import csv
import json
import random
from collections import defaultdict
from pathlib import Path

SEED = 20260729
# all generation-scored arms, spanning the full answer-quality range (closed floor ->
# ceiling) so the judge is validated across every answer style, not just the strong ones.
AUDIT_ARMS = ["C0", "R0", "R1", "A1", "A2", "A3", "A5", "A4B", "C1"]


def _load(p):
    return [json.loads(l) for l in Path(p).read_text("utf-8").splitlines() if l.strip()]


def _index(rows, key="id"):
    return {r[key]: r for r in rows}


def sample(run: Path, n: int):
    ev = _index(_load("eval/eval_set.generated.jsonl"))
    out_dir = run / "judge_audit"
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(SEED)

    # pool candidate (arm,item) verdicts, joined with the answer text
    pool = []
    for arm in AUDIT_ARMS:
        verd = _index(_load(run / "verdicts" / f"{arm}.jsonl"))
        ans = _index(_load(run / "answers" / f"{arm}.jsonl"))
        for qid, v in verd.items():
            if "correctness" not in v or qid not in ans or qid not in ev:
                continue
            pool.append({
                "id": qid, "arm": arm,
                "question_type": ev[qid]["question_type"],
                "question": ev[qid]["question"],
                "reference_answer": ev[qid]["reference_answer"],
                "model_answer": ans[qid]["answer"],
                "judge_correct": 1 if v["correctness"]["correct"] else 0,
                "judge_score": int(v["correctness"].get("score", 0)),
                "judge_rationale": v["correctness"].get("rationale", ""),
            })

    # COVERAGE-FIRST stratification: equal allocation per question_type, and within each
    # type balance the judge's correct/incorrect classes while spreading across arms.
    by_type = defaultdict(list)
    for r in pool:
        by_type[r["question_type"]].append(r)
    types = sorted(by_type)
    per_type = n // len(types)
    picked, seen = [], set()

    def take(cands, target):
        """Round-robin over arms to spread answer styles; avoid duplicate (id,arm)."""
        by_arm = defaultdict(list)
        for r in cands:
            by_arm[r["arm"]].append(r)
        for a in by_arm:
            rng.shuffle(by_arm[a])
        arms_order = list(by_arm)
        rng.shuffle(arms_order)
        got = 0
        while got < target and any(by_arm.values()):
            for a in arms_order:
                if by_arm[a]:
                    r = by_arm[a].pop()
                    if (r["id"], r["arm"]) in seen:
                        continue
                    seen.add((r["id"], r["arm"]))
                    picked.append(r)
                    got += 1
                    if got >= target:
                        break
        return got

    for t in types:
        rows = by_type[t]
        pos = [r for r in rows if r["judge_correct"] == 1]
        neg = [r for r in rows if r["judge_correct"] == 0]
        half = per_type // 2
        got = take(pos, half) + take(neg, per_type - half)
        # if one class is short in this type, top up from the other so the type still hits quota
        if got < per_type:
            take(pos + neg, per_type - got)
    # top up to exactly n (rounding remainder) from anything unused
    if len(picked) < n:
        rest = [r for r in pool if (r["id"], r["arm"]) not in seen]
        rng.shuffle(rest)
        for r in rest[: n - len(picked)]:
            seen.add((r["id"], r["arm"]))
            picked.append(r)
    rng.shuffle(picked)

    # write blind sheet (no judge columns) + hidden key
    blind = out_dir / "audit_sheet_blind.csv"
    key = out_dir / "audit_key.csv"
    with blind.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["audit_id", "question_type", "question", "reference_answer",
                    "model_answer", "HUMAN_correct", "HUMAN_score", "notes"])
        for i, r in enumerate(picked):
            w.writerow([i, r["question_type"], r["question"], r["reference_answer"],
                        r["model_answer"], "", "", ""])
    with key.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["audit_id", "id", "arm", "judge_correct", "judge_score", "judge_rationale"])
        for i, r in enumerate(picked):
            w.writerow([i, r["id"], r["arm"], r["judge_correct"], r["judge_score"], r["judge_rationale"]])

    nc = sum(r["judge_correct"] for r in picked)
    print(f"✓ wrote {len(picked)}-item BLIND audit sheet -> {blind}")
    print(f"  (judge says: {nc} correct / {len(picked)-nc} incorrect; hidden key -> {key})")
    print(f"\nNEXT: open {blind}, fill HUMAN_correct (1=correct, 0=incorrect) and HUMAN_score (1-5),")
    print(f"then run:  .venv/bin/python -m analysis.judge_audit --run {run} --score")


def _kappa(a, b):
    """Cohen's kappa for two binary label lists."""
    n = len(a)
    po = sum(1 for x, y in zip(a, b) if x == y) / n
    # expected agreement
    pe = 0.0
    for c in (0, 1):
        pe += (a.count(c) / n) * (b.count(c) / n)
    return (po - pe) / (1 - pe) if pe != 1 else 1.0, po


def _pearson(x, y):
    n = len(x)
    mx, my = sum(x) / n, sum(y) / n
    num = sum((xi - mx) * (yi - my) for xi, yi in zip(x, y))
    dx = sum((xi - mx) ** 2 for xi in x) ** 0.5
    dy = sum((yi - my) ** 2 for yi in y) ** 0.5
    return num / (dx * dy) if dx and dy else float("nan")


def _rank(v):
    order = sorted(range(len(v)), key=lambda i: v[i])
    r = [0.0] * len(v)
    i = 0
    while i < len(v):
        j = i
        while j + 1 < len(v) and v[order[j + 1]] == v[order[i]]:
            j += 1
        avg = (i + j) / 2 + 1
        for k in range(i, j + 1):
            r[order[k]] = avg
        i = j + 1
    return r


def score(run: Path):
    out_dir = run / "judge_audit"
    blind = {r["audit_id"]: r for r in csv.DictReader((out_dir / "audit_sheet_blind.csv").open(encoding="utf-8"))}
    key = {r["audit_id"]: r for r in csv.DictReader((out_dir / "audit_key.csv").open(encoding="utf-8"))}

    recs, skipped = [], 0   # each: (question_type, judge_bin, human_bin, judge_score, human_score|None)
    for aid, b in blind.items():
        hc, hs = b.get("HUMAN_correct", "").strip(), b.get("HUMAN_score", "").strip()
        if hc == "" or aid not in key:
            skipped += 1
            continue
        j = int(key[aid]["judge_correct"])
        h = 1 if hc in ("1", "true", "True", "yes") else 0
        js = int(key[aid]["judge_score"])
        hsv = int(float(hs)) if hs else None
        recs.append((b.get("question_type", "?"), j, h, js, hsv))
    n = len(recs)
    if n == 0:
        print("No HUMAN_correct labels filled yet — nothing to score.")
        return

    def _interp(k):
        return ("almost perfect" if k >= .81 else "substantial" if k >= .61 else
                "moderate" if k >= .41 else "fair" if k >= .21 else "slight/poor")

    def _report(rs, label, indent=""):
        hj = [r[1] for r in rs]
        hh = [r[2] for r in rs]
        kappa, po = _kappa(hj, hh)
        sj = [r[3] for r in rs if r[4] is not None]
        sh = [r[4] for r in rs if r[4] is not None]
        # kappa is undefined when either rater used a single class -> fall back to agreement
        one_class = len(set(hj)) < 2 or len(set(hh)) < 2
        ktxt = "n/a (one class)" if one_class else f"{kappa:.3f} ({_interp(kappa)})"
        extra = ""
        if sj and len(set(sj)) > 1 and len(set(sh)) > 1:
            extra = f"  | score Pearson={_pearson(sj, sh):.2f} Spearman={_pearson(_rank(sj), _rank(sh)):.2f}"
        print(f"{indent}{label:<14} n={len(rs):<3} agree={po:.3f}  kappa={ktxt}{extra}")

    print(f"Judge validation — {n} human-labelled items ({skipped} unlabelled skipped)\n")
    _report(recs, "OVERALL")
    print("\n  per question_type:")
    for t in sorted(set(r[0] for r in recs)):
        _report([r for r in recs if r[0] == t], t, indent="    ")

    # overall confusion + disagreements to inspect
    cm = [[0, 0], [0, 0]]
    for _, j, h, _, _ in recs:
        cm[j][h] += 1
    print("\n  confusion (rows=judge, cols=human):")
    print(f"                 human=incorrect  human=correct")
    print(f"    judge=incorrect     {cm[0][0]:>6}          {cm[0][1]:>6}")
    print(f"    judge=correct       {cm[1][0]:>6}          {cm[1][1]:>6}")
    dis = cm[0][1] + cm[1][0]
    print(f"\n  disagreements: {dis}/{n} ({dis/n:.0%}) — false-pos (judge✓/human✗)={cm[1][0]}, "
          f"false-neg (judge✗/human✓)={cm[0][1]}")
    bad = [aid for aid, b in blind.items() if aid in key and b.get("HUMAN_correct", "").strip() != ""
           and int(key[aid]["judge_correct"]) != (1 if b["HUMAN_correct"].strip() in ("1", "true", "True", "yes") else 0)]
    if bad:
        print(f"  disagreement audit_ids: {', '.join(sorted(bad, key=lambda x: int(x)))}")


def export_xlsx(run: Path):
    """Convert the blind CSV into an .xlsx with dropdown data validation on the two
    HUMAN columns (0/1 and 1-5). Opens in Excel; keeps the dropdowns when imported to
    Google Sheets. Read-only source columns are wrapped/sized; HUMAN columns highlighted."""
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    src = run / "judge_audit" / "audit_sheet_blind.csv"
    rows = list(csv.DictReader(src.open(encoding="utf-8")))
    if not rows:
        raise SystemExit(f"{src} is empty — run --sample first")
    headers = list(rows[0].keys())
    wb = Workbook()
    ws = wb.active
    ws.title = "judge_audit"

    hdr_fill = PatternFill("solid", fgColor="34495E")
    human_fill = PatternFill("solid", fgColor="FCF3CF")   # highlight the columns to fill
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([r[h] for h in headers])

    widths = {"audit_id": 8, "question_type": 13, "question": 60, "reference_answer": 60,
              "model_answer": 60, "HUMAN_correct": 15, "HUMAN_score": 13, "notes": 24}
    col = {h: get_column_letter(i + 1) for i, h in enumerate(headers)}
    for h, w in widths.items():
        if h in col:
            ws.column_dimensions[col[h]].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for h in ("question", "reference_answer", "model_answer", "notes"):
        if h in col:
            for cell in ws[col[h]][1:]:
                cell.alignment = wrap

    n = len(rows)
    dvs = {"HUMAN_correct": '"0,1"', "HUMAN_score": '"1,2,3,4,5"'}
    for h, f in dvs.items():
        if h not in col:
            continue
        dv = DataValidation(type="list", formula1=f, allow_blank=True, showDropDown=False)
        dv.error = "เลือกจากรายการเท่านั้น"
        dv.errorTitle = "ค่าไม่ถูกต้อง"
        ws.add_data_validation(dv)
        dv.add(f"{col[h]}2:{col[h]}{n + 1}")
        for cell in ws[col[h]][1:]:
            cell.fill = human_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"       # keep header visible while scrolling
    out = run / "judge_audit" / "audit_sheet.xlsx"
    wb.save(out)
    print(f"✓ wrote {out}  ({n} rows; dropdowns on HUMAN_correct[0/1] + HUMAN_score[1-5])")
    print("  Excel: fill the yellow columns. Google Sheets: File→Import→Upload (keeps dropdowns).")
    print("  When done, export back to CSV over audit_sheet_blind.csv, then run --score.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", type=Path, default=Path("runs/2026-08-07"))
    ap.add_argument("--sample", action="store_true")
    ap.add_argument("--score", action="store_true")
    ap.add_argument("--xlsx", action="store_true", help="export blind sheet to .xlsx with dropdowns")
    ap.add_argument("--n", type=int, default=80)
    args = ap.parse_args()
    if args.sample:
        sample(args.run, args.n)
    elif args.xlsx:
        export_xlsx(args.run)
    elif args.score:
        score(args.run)
    else:
        ap.error("choose --sample, --xlsx, or --score")


if __name__ == "__main__":
    main()
