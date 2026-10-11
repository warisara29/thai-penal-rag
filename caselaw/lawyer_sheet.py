"""Build a lawyer-facing QC worksheet for the Penal-Code citation extraction (task F review).

Unlike qc_extract.py (a dev parser-vs-parser diff), this sheet asks a Thai lawyer the only
question that matters: *for this case, are the sections we extracted correct?* Each row shows
the case, the sections our PRODUCTION path extracted (the dataset's curated field, normalized +
mapped to the corpus), and a focused excerpt of the judgment around every มาตรา / พิพากษา
mention — so the reviewer can confirm without reading the whole opinion. Three blank columns
collect the verdict: all-correct?, sections missed, sections over-included.

Cases where an independent full-text regex surfaced a *substantive* (non-boilerplate) section
the curated field lacks, or where a cited section failed to map, are flagged ต้องดูพิเศษ and
sorted to the top so limited lawyer time lands on the highest-value rows first.

Usage:
    python -m caselaw.lawyer_sheet --val data/caselaw/raw/val.jsonl \
        --mapped data/caselaw/penal_citations.mapped.jsonl --n 100 --seed 20261006
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import re

from .qc_extract import _HDR, _load_text_by_caseid, _sortkey
from .sections import extract_from_text

# Sentencing / general-part mechanics that routinely appear in the body but are NOT what a case
# is "about" — used only to decide whether a regex-only extra is worth a lawyer's attention.
_BOILERPLATE = {
    "32", "33", "38", "56", "58", "59", "78", "80", "83", "84", "85", "86",
    "90", "91", "92", "93", "96", "371",
}


def _excerpt(text: str, max_chars: int = 2200) -> str:
    """Judgment header + merged ±windows around every มาตรา/พิพากษา mention, capped."""
    t = re.sub(r"[ \t]+", " ", text).strip()
    spans = [(0, 380)]  # always include the charge/head-note intro
    for m in re.finditer(r"มาตรา|พิพากษา|ม\.\s*\d", t):
        spans.append((max(0, m.start() - 60), m.start() + 90))
    spans.sort()
    merged: list[list[int]] = []
    for a, b in spans:
        if merged and a <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    out, total = [], 0
    for a, b in merged:
        piece = t[a:b]
        out.append(piece)
        total += len(piece)
        if total >= max_chars:
            break
    return " … ".join(out)


def _fmt_sections(keys: list[str]) -> str:
    return "  ".join(f"มาตรา {k}" for k in sorted(keys, key=_sortkey))


def run(val_path: str, mapped_path: str, n: int, seed: int, out_base: str) -> dict:
    text_by_id = _load_text_by_caseid(val_path)
    cases: dict[str, dict] = {}
    with open(mapped_path, encoding="utf-8") as f:
        for line in f:
            d = json.loads(line)
            if d["case_id"] in text_by_id and d["case_id"] not in cases:
                cases[d["case_id"]] = d
    sample = random.Random(seed).sample(sorted(cases), min(n, len(cases)))

    rows = []
    for cid in sample:
        d = cases[cid]
        text = text_by_id[cid]
        curated = set(d["sections"])                      # production extraction (keys)
        regex = set(extract_from_text(text))
        # regex-only extras that are NOT sentencing boilerplate -> a possible missed citation
        substantive_extra = sorted(
            (regex - curated) - _BOILERPLATE, key=_sortkey)
        flag = bool(d.get("unmapped")) or bool(substantive_extra)
        rows.append({
            "_flag": flag,
            "ต้องดูพิเศษ": "★" if flag else "",
            "เลขคดี": cid,
            "ปี_พศ": d.get("year_be") or "",
            "มาตราที่ระบบดึงได้": _fmt_sections(d["sections"]),
            "มาตราที่ยังแมปไม่ได้": _fmt_sections(d.get("unmapped", [])),
            "regex_พบเพิ่ม_ควรเช็ก": _fmt_sections(substantive_extra),
            "ข้อความคำพิพากษา_ช่วงที่อ้างมาตรา": _excerpt(text),
            "ถูกทั้งหมด": "",          # lawyer: ✓ / ✗
            "มาตราที่ขาด": "",         # lawyer fills
            "มาตราที่เกิน": "",        # lawyer fills
            "หมายเหตุ": "",
        })
    rows.sort(key=lambda r: (not r["_flag"]))  # flagged first
    for r in rows:
        r.pop("_flag")

    csv_path = f"{out_base}.csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:  # BOM: Excel reads Thai
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    xlsx_path = _write_xlsx(rows, f"{out_base}.xlsx")
    return {"sampled": len(rows), "flagged": sum(bool(r["ต้องดูพิเศษ"]) for r in rows),
            "csv": csv_path, "xlsx": xlsx_path, "joinable": len(cases)}


def _write_xlsx(rows: list[dict], out: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    headers = list(rows[0].keys())
    wb = Workbook()
    ws = wb.active
    ws.title = "lawyer_qc"
    ws.append(headers)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    human_fill = PatternFill("solid", fgColor="FCF3CF")   # columns the lawyer fills
    flag_fill = PatternFill("solid", fgColor="FADBD8")    # flagged rows
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r[h] for h in headers])

    widths = {"ต้องดูพิเศษ": 9, "เลขคดี": 11, "ปี_พศ": 8, "มาตราที่ระบบดึงได้": 34,
              "มาตราที่ยังแมปไม่ได้": 16, "regex_พบเพิ่ม_ควรเช็ก": 18,
              "ข้อความคำพิพากษา_ช่วงที่อ้างมาตรา": 90, "ถูกทั้งหมด": 11,
              "มาตราที่ขาด": 16, "มาตราที่เกิน": 16, "หมายเหตุ": 24}
    col = {h: get_column_letter(i + 1) for i, h in enumerate(headers)}
    for h, w in widths.items():
        if h in col:
            ws.column_dimensions[col[h]].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for h in ("มาตราที่ระบบดึงได้", "ข้อความคำพิพากษา_ช่วงที่อ้างมาตรา", "หมายเหตุ",
              "มาตราที่ยังแมปไม่ได้", "regex_พบเพิ่ม_ควรเช็ก"):
        for cell in ws[col[h]][1:]:
            cell.alignment = wrap

    n = len(rows)
    # dropdown ✓/✗ on ถูกทั้งหมด
    dv = DataValidation(type="list", formula1='"✓,✗"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col['ถูกทั้งหมด']}2:{col['ถูกทั้งหมด']}{n + 1}")
    for h in ("ถูกทั้งหมด", "มาตราที่ขาด", "มาตราที่เกิน", "หมายเหตุ"):
        for cell in ws[col[h]][1:]:
            cell.fill = human_fill
            if h == "ถูกทั้งหมด":
                cell.alignment = Alignment(horizontal="center", vertical="center")
    # tint flagged rows (col A has ★)
    for i, r in enumerate(rows, start=2):
        if r["ต้องดูพิเศษ"]:
            ws.cell(i, 1).fill = flag_fill
    ws.freeze_panes = "C2"   # keep ★ + case id + header visible
    wb.save(out)
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--val", default="data/caselaw/raw/val.jsonl")
    ap.add_argument("--mapped", default="data/caselaw/penal_citations.mapped.jsonl")
    ap.add_argument("--n", type=int, default=100)
    ap.add_argument("--seed", type=int, default=20261006)
    ap.add_argument("--out", default="data/caselaw/lawyer_qc_sheet")
    args = ap.parse_args()
    s = run(args.val, args.mapped, args.n, args.seed, args.out)
    print(f"joinable cases: {s['joinable']}  |  sampled: {s['sampled']}  "
          f"|  flagged ต้องดูพิเศษ: {s['flagged']}")
    print(f"  {s['xlsx']}   (open in Excel / Google Sheets — ✓/✗ dropdown on ถูกทั้งหมด)")
    print(f"  {s['csv']}")


if __name__ == "__main__":
    main()
